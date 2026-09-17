import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

doc = Document()

# Налаштування полів (Верхнє, Нижнє, Праве: 2 см; Ліве: 2.5 см)
for section in doc.sections:
    section.top_margin = Inches(0.79)
    section.bottom_margin = Inches(0.79)
    section.left_margin = Inches(0.98)
    section.right_margin = Inches(0.79)

def p(text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, italic=False, size=11, space_after=4, space_before=0, color=None):
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(space_before)
    r = paragraph.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color:
        r.font.color.rgb = color
    return paragraph

# --- ТИТУЛЬНА СТОРІНКА ---
p("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ\nНАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ БІОРЕСУРСІВ І ПРИРОДОКОРИСТУВАННЯ УКРАЇНИ\nФакультет інформаційних технологій", WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12)

doc.add_paragraph().paragraph_format.space_before = Pt(72)
p("ЗВІТ\nз лабораторної роботи №4\nна тему: «Модулі і пакети»", WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16)

doc.add_paragraph().paragraph_format.space_before = Pt(96)
p("Виконав:\nстудент групи ІТ-31\nСніжко І.Б.\n\nПеревірив:\nвикладач _____________", WD_ALIGN_PARAGRAPH.RIGHT, size=12)

doc.add_paragraph().paragraph_format.space_before = Pt(72)
p("Київ – 2026", WD_ALIGN_PARAGRAPH.CENTER, size=12)

doc.add_page_break()

# --- ТЕКСТ ЗАВДАННЯ ---
p("1. Мета та текст завдання", bold=True, size=14, space_after=6)
task_text = (
    "Мета роботи: Навчитися створювати та використовувати власні модулі і пакети в Python, "
    "працювати з віртуальним оточенням, конфігураційними файлами та зовнішніми бібліотеками перекладу.\n\n"
    "Основні завдання:\n"
    "1. Створити віртуальне оточення (ім'я оточення: Snizhko).\n"
    "2. Створити пакет my_package з трьох модулів (googletrans 4.0.2, googletrans 3.1.0a0, deep-translator + langdetect).\n"
    "3. Реалізувати змінні NAME та AUTHOR в файлі __init__.py.\n"
    "4. Створити файли gtrans4.py, gtrans3.py, deeptr.py для демонстрації роботи модулів.\n"
    "5. Створити файл filetr.py для перекладу тексту з файлу відповідно до конфігураційного файлу (config.json).\n"
    "6. Створити файли requirements.txt та .gitignore.\n"
    "7. Завантажити проєкт на GitHub."
)
p(task_text, size=11, space_after=12)

# --- КОД ГОЛОВНОЇ ПРОГРАМИ ---
p("2. Код головної програми (filetr.py)", bold=True, size=14, space_after=6)

try:
    with open("filetr.py", "r", encoding="utf-8") as f:
        code_content = f.read()
except Exception:
    code_content = "# Файл filetr.py не знайдено в поточній папці"

table = doc.add_table(rows=1, cols=1)
table.autofit = False
cell = table.cell(0, 0)
set_cell_background(cell, "F4F4F4")
set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

cp = cell.paragraphs[0]
cp.paragraph_format.space_after = Pt(2)
cp.paragraph_format.space_before = Pt(2)
run = cp.add_run(code_content)
run.font.name = 'Consolas'
run.font.size = Pt(9.5)
run.font.color.rgb = RGBColor(0x1E, 0x1E, 0x1E)

# --- СКРІНШОТИ ТА ПОСИЛАННЯ ---
doc.add_paragraph().paragraph_format.space_before = Pt(12)
p("3. Результати виконання та скріншоти", bold=True, size=14, space_after=6)

p("Нижче наведено підтвердження успішного виконання лабораторної роботи у віртуальному оточенні Snizhko:")

p("[Вставити скріншот VS Code зі структурою проєкту та блакитним рядком стану]", italic=True, space_after=4, color=RGBColor(0x7F, 0x7F, 0x7F))
p("[Вставити скріншот виконаної програми filetr.py у терміналі Snizhko]", italic=True, space_after=4, color=RGBColor(0x7F, 0x7F, 0x7F))
p("[Вставити скріншот повідомлення про помилку версії Python 3.13 для gtrans3.py]", italic=True, space_after=12, color=RGBColor(0x7F, 0x7F, 0x7F))

p("Посилання на GitHub-репозиторій проєкту:", bold=True, size=11, space_after=2)
p("https://github.com/vanya-snizhko/lab_project", size=11, color=RGBColor(0x00, 0x00, 0xEE))

output_filename = "Звіт_Лабораторна_4_Сніжко.docx"
doc.save(output_filename)
print(f"✅ Звіт успішно створено без помилок кодування: {output_filename}")