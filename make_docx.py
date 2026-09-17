import docx
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
for s in doc.sections:
    s.top_margin = Inches(0.79)
    s.bottom_margin = Inches(0.79)
    s.left_margin = Inches(0.98)
    s.right_margin = Inches(0.79)

def p(text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, italic=False, size=11, space_after=4):
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    r = paragraph.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    return paragraph

p("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ\n Університет економіки та Права "КРОК""\n, WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12)
doc.add_paragraph().paragraph_format.space_before = Pt(60)
p("ЗВІТ\nз лабораторної роботи №4\nна тему: «Модулі і пакети»", WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16)
doc.add_paragraph().paragraph_format.space_before = Pt(80)
p("Виконав:\nстудент групи ІТ-31\nСніжко І.Б.\n", WD_ALIGN_PARAGRAPH.RIGHT, size=12)
doc.add_paragraph().paragraph_format.space_before = Pt(60)
p("Київ – 2026", WD_ALIGN_PARAGRAPH.CENTER, size=12)

doc.add_page_break()

p("1. Текст завдання", bold=True, size=14, space_after=6)
p("Створити віртуальне оточення, створити власний пакет my_package з трьох модулів для перекладу тексту, реалізувати обробку конфігураційних файлів у filetr.py та завантажити проєкт на GitHub.")

p("2. Код головної програми (filetr.py)", bold=True, size=14, space_after=6)
with open("filetr.py", "r", encoding="utf-8") as f:
    code_text = f.read()
p(code_text, size=9.5)

p("3. Результати та скріншоти", bold=True, size=14, space_after=6)
p("• Програма виконана у віртуальному оточенні Snizhko.")
p("• Всі модулі пакета успішно перекладають текст та визначають мову.")
p("• Посилання на GitHub: https://github.com/SnizhkoIvan/Zavdannya_3")
